# -*- coding: utf-8 -*-
"""분개장 읽기: 시트/헤더 자동 연결, 합계·반복헤더·빈 행 제외."""
from __future__ import annotations

import datetime as _dt
import re
import unicodedata
from dataclasses import dataclass, field
from typing import Any, Iterable

import openpyxl

FIELD_KEYS = (
    "voucher_id", "date", "voucher_no", "dc",
    "account", "debit", "credit", "counterparty", "description",
)


class JournalReadError(Exception):
    """분개장을 읽을 수 없을 때. '후보 0건'과 반드시 구분한다."""


@dataclass
class JournalRow:
    sheet: str
    excel_row: int
    voucher_id: str = ""
    date: _dt.date | None = None
    voucher_no: str = ""
    dc: str = ""
    account: str = ""
    debit: float = 0.0
    credit: float = 0.0
    counterparty: str = ""
    description: str = ""

    def text_of(self, key: str) -> str:
        return getattr(self, key, "") or ""


@dataclass
class JournalData:
    path: str
    sheet: str
    profile_name: str
    header_row: int
    column_map: dict[str, int]
    rows: list[JournalRow] = field(default_factory=list)
    skipped: dict[str, int] = field(default_factory=dict)
    date_min: _dt.date | None = None
    date_max: _dt.date | None = None
    missing_fields: list[str] = field(default_factory=list)


def _norm_header(v: Any) -> str:
    if v is None:
        return ""
    s = unicodedata.normalize("NFKC", str(v))
    return re.sub(r"\s+", "", s).lower()


def _to_date(v: Any) -> _dt.date | None:
    if v is None:
        return None
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    s = unicodedata.normalize("NFKC", str(v)).strip()
    m = re.match(r"^(\d{4})\D(\d{1,2})\D(\d{1,2})", s)
    if not m:
        m = re.match(r"^(\d{4})(\d{2})(\d{2})$", s)
    if m:
        try:
            return _dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        except ValueError:
            return None
    return None


def _to_amount(v: Any) -> float:
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[,\s원]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return 0.0


def _to_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def _match_profile(ws, profile: dict[str, Any]) -> tuple[int, dict[str, int]] | None:
    """헤더 행과 열 위치를 찾는다. 실패하면 None."""
    search_rows = int(profile.get("header_search_rows", 15))
    cols_cfg: dict[str, list[str]] = profile.get("columns", {})
    wanted = {k: [_norm_header(n) for n in names] for k, names in cols_cfg.items()}

    for r, row in enumerate(ws.iter_rows(min_row=1, max_row=search_rows, values_only=True), start=1):
        cells = [_norm_header(v) for v in row]
        cmap: dict[str, int] = {}
        for key, names in wanted.items():
            for idx, cell in enumerate(cells, start=1):
                if cell and cell in names:
                    cmap[key] = idx
                    break
        required = profile.get("required", ["account", "counterparty"])
        if all(k in cmap for k in required) and len(cmap) >= 3:
            return r, cmap
    return None


def read_journal(path: str, profiles_cfg: dict[str, Any],
                 profile_name: str | None = None) -> JournalData:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise JournalReadError(f"분개장 파일을 열 수 없습니다: {exc}") from exc

    profiles: dict[str, Any] = profiles_cfg.get("profiles", {})
    if profile_name:
        if profile_name not in profiles:
            wb.close()
            raise JournalReadError(f"입력 프로파일 '{profile_name}'이(가) 설정에 없습니다.")
        order = [profile_name]
    else:
        default = profiles_cfg.get("default_profile")
        order = ([default] if default in profiles else []) + [p for p in profiles if p != default]

    best: tuple[int, str, dict[str, Any], Any, int, dict[str, int]] | None = None
    for pname in order:
        profile = profiles[pname]
        cands = [s for s in profile.get("sheet_candidates", []) if s in wb.sheetnames]
        sheets = cands + [s for s in wb.sheetnames if s not in cands]
        for sname in sheets:
            ws = wb[sname]
            found = _match_profile(ws, profile)
            if found:
                hrow, cmap = found
                score = len(cmap) + (10 if sname in cands else 0)
                if best is None or score > best[0]:
                    best = (score, pname, profile, sname, hrow, cmap)
        if best and best[1] == pname and best[0] >= 10:
            break

    if best is None:
        wb.close()
        raise JournalReadError(
            "분개장에서 헤더(계정과목·거래처 등)를 찾지 못했습니다. "
            "config/input_profiles.json의 열 이름을 확인하세요."
        )

    _, pname, profile, sname, header_row, cmap = best
    ws = wb[sname]
    rules = profiles_cfg.get("skip_row_rules", {})
    summary_tokens = rules.get("summary_tokens", ["합계", "소계", "총계"])
    require_date = bool(rules.get("require_parsable_date", True)) and "date" in cmap

    header_norm = {_norm_header(n) for names in profile.get("columns", {}).values() for n in names}

    rows: list[JournalRow] = []
    skipped = {"빈행": 0, "합계·소계행": 0, "반복헤더": 0, "날짜없음": 0}
    dmin = dmax = None

    for excel_row, raw in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                    start=header_row + 1):
        if raw is None:
            skipped["빈행"] += 1
            continue
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in raw):
            skipped["빈행"] += 1
            continue

        def cell(key: str) -> Any:
            i = cmap.get(key)
            if i is None or i > len(raw):
                return None
            return raw[i - 1]

        joined = " ".join(_to_text(v) for v in raw if v is not None)
        if any(tok in joined for tok in summary_tokens) and not _to_date(cell("date")):
            skipped["합계·소계행"] += 1
            continue
        if rules.get("drop_repeated_header", True):
            vals = {_norm_header(cell(k)) for k in ("account", "counterparty", "description")}
            vals.discard("")
            if vals and vals <= header_norm:
                skipped["반복헤더"] += 1
                continue

        d = _to_date(cell("date"))
        if require_date and d is None:
            skipped["날짜없음"] += 1
            continue

        jr = JournalRow(
            sheet=sname,
            excel_row=excel_row,
            voucher_id=_to_text(cell("voucher_id")),
            date=d,
            voucher_no=_to_text(cell("voucher_no")),
            dc=_to_text(cell("dc")),
            account=_to_text(cell("account")),
            debit=_to_amount(cell("debit")),
            credit=_to_amount(cell("credit")),
            counterparty=_to_text(cell("counterparty")),
            description=_to_text(cell("description")),
        )
        if not jr.voucher_id:
            jr.voucher_id = f"{jr.date or ''}-{jr.voucher_no}" if (jr.date or jr.voucher_no) else f"ROW{excel_row}"
        if d:
            dmin = d if dmin is None or d < dmin else dmin
            dmax = d if dmax is None or d > dmax else dmax
        rows.append(jr)

    wb.close()

    missing = [k for k in FIELD_KEYS if k not in cmap]
    if not rows:
        raise JournalReadError(
            f"'{sname}' 시트에서 거래행을 한 건도 읽지 못했습니다(제외 내역: {skipped}). "
            "검색 결과 0건과 다른 상황이므로 저장을 진행하지 않습니다."
        )

    return JournalData(
        path=path, sheet=sname, profile_name=pname, header_row=header_row,
        column_map=cmap, rows=rows, skipped=skipped,
        date_min=dmin, date_max=dmax, missing_fields=missing,
    )
